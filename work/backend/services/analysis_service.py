"""
Analysis Service
----------------
Takes a user's uploaded transaction history (a CSV) and runs the full
pipeline against it: clean -> engineer features -> score anomalies ->
detect subscriptions -> spending breakdown (by category AND by specific
merchant/shop) -> a dashboard-ready summary. This is what backs
POST /upload/analyze.

Expected upload CSV columns (case-insensitive, flexible):
  transaction_id (optional, generated if missing), timestamp/date,
  amount, merchant, category (optional, inferred as "Other" if missing),
  payment_mode (optional)
"""

import io
import os
import sys
import uuid
import joblib
import numpy as np
import pandas as pd
from sklearn.ensemble import IsolationForest
from sklearn.preprocessing import StandardScaler

ML_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "ml")
sys.path.insert(0, ML_DIR)

from features import build_user_spending_features, build_recurring_features   # noqa: E402
from anomaly_model import FEATURE_COLS, explain_row, MODEL_PATH, SCALER_PATH   # noqa: E402
from insights_engine import detect_period, generate_insights                    # noqa: E402
from subscription_model import detect_subscriptions                            # noqa: E402
from categorizer import fill_categories                                         # noqa: E402

REQUIRED_COLS = {"amount", "merchant"}
MIN_TXNS_FOR_PERSONAL_MODEL = 20
RULE_MOD_Z = 3.5   # robust z-score above which an amount is flagged as unusually high


def _normalize_columns(df):
    df.columns = [c.strip().lower().replace(" ", "_") for c in df.columns]

    # amount: common export names
    amount_aliases = ["amount_(inr)", "amount_inr", "txn_amount", "transaction_amount",
                       "debit_amount", "amount_(rs.)", "value", "amt"]
    # merchant: real bank/UPI exports rarely say "merchant" literally
    merchant_aliases = ["merchant_name", "payee", "payee_name", "beneficiary", "beneficiary_name",
                         "receiver", "receiver_name", "to", "party", "party_name", "narration",
                         "description", "remarks", "particulars", "transaction_details", "vpa",
                         "counterparty", "paid_to_/_received_from", "paid_to", "received_from",
                         "paid_to_received_from"]
    # date/time
    date_aliases = ["date", "txn_date", "transaction_date", "value_date", "posting_date", "datetime",
                     "date_&_time", "date_time", "date_and_time"]
    # category (NOTE: deliberately does NOT include "type" — on real bank/UPI
    # exports "Type" almost always means Debit/Credit, not a spending
    # category, and mapping it to category would silently corrupt the
    # category breakdown)
    category_aliases = ["merchant_category", "txn_category"]

    rename_map = {}
    for alias in amount_aliases:
        if alias in df.columns:
            rename_map[alias] = "amount"
            break
    for alias in merchant_aliases:
        if alias in df.columns:
            rename_map[alias] = "merchant"
            break
    for alias in date_aliases:
        if alias in df.columns:
            rename_map[alias] = "timestamp"
            break
    for alias in category_aliases:
        if alias in df.columns:
            rename_map[alias] = "category"
            break

    df = df.rename(columns=rename_map)
    return df


def _parse_amount(val):
    """Clean a currency-formatted amount into a float, or NaN if it isn't
    really a number. Handles real-world export quirks: '₹1,450.00',
    'Rs. 820', 'INR 299', a leading '+' on credits, stray '-', and a bare
    '+' with no digits at all (seen on malformed rows — treated as
    unusable rather than crashing the whole upload)."""
    if pd.isna(val):
        return np.nan
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    s = s.replace("₹", "").replace("Rs.", "").replace("Rs", "").replace("INR", "")
    s = s.replace(",", "").replace("+", "").strip()
    if s in ("", "-", "."):
        return np.nan
    try:
        return float(s)
    except ValueError:
        return np.nan


def _is_short_lowercase_fragment(s):
    """Heuristic for reassembling a text cell that Excel wrapped across
    rows: a short, all-lowercase fragment ('ket', 'nt', 'd', 'k') is almost
    always the tail of the previous word ('Supermar'+'ket'->'Supermarket'),
    while a longer or capitalized fragment ('from', 'Friend') is a
    separate word and should get a space before it. Not perfect, but good
    enough to turn unreadable fragments into a readable merchant name."""
    return bool(s) and len(s) <= 3 and s.isalpha() and s.islower()


def _join_text_fragments(values):
    """Reassemble a wrapped text cell from its ordered row fragments."""
    out = ""
    for v in values:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            continue
        frag = str(v).replace("\n", " ").strip()
        if not frag:
            continue
        if not out:
            out = frag
        elif _is_short_lowercase_fragment(frag):
            out += frag
        else:
            out += " " + frag
    return out.strip()


def _reconstruct_wrapped_rows(df):
    """Some spreadsheet/statement exports wrap a single logical transaction
    across several physical rows (long merchant names, multi-line dates),
    leaving most cells in the continuation rows blank. This is common in
    .xlsx exports where the original cells had word-wrap enabled.

    We detect this by looking for a reliable "this row starts a new
    transaction" marker — a 'status' or 'type' (Debit/Credit) column that is
    only ever populated on the first physical row of each transaction. If
    such a marker exists and some rows are blank on it, we group the
    continuation rows into their parent transaction and reassemble each
    column instead of silently losing data.

    If no such marker is found (or every row already has it), the frame is
    returned unchanged — this only activates for genuinely fragmented
    files."""
    marker_col = None
    for candidate in ("status", "type", "txn_type"):
        if candidate in df.columns and df[candidate].notna().any():
            marker_col = candidate
            break

    if marker_col is None:
        return df

    is_start = df[marker_col].notna()
    if is_start.all() or not is_start.any():
        return df  # nothing fragmented, or no usable marker at all

    # Drop any stray rows before the first real transaction starts.
    first_start = is_start.idxmax()
    df = df.loc[first_start:].reset_index(drop=True)
    is_start = df[marker_col].notna()
    group_id = is_start.cumsum()

    text_cols = [c for c in ("timestamp", "merchant", "payment_mode") if c in df.columns]
    passthrough_cols = [c for c in df.columns if c not in text_cols + ["amount"]]

    rows = []
    for _, group in df.groupby(group_id):
        rec = {}
        for col in text_cols:
            rec[col] = _join_text_fragments(group[col].tolist())
        if "amount" in df.columns:
            rec["amount"] = np.nan
            for v in group["amount"].tolist():
                parsed = _parse_amount(v)
                if not pd.isna(parsed):
                    rec["amount"] = parsed
                    break
        for col in passthrough_cols:
            non_null = group[col].dropna()
            rec[col] = non_null.iloc[0] if len(non_null) else np.nan
        rows.append(rec)

    return pd.DataFrame(rows)


def _read_json_transactions(file_bytes):
    """Accepts a JSON export in any of the shapes people actually send:
    a top-level list of transaction objects, {"transactions": [...]},
    {"data": [...]}, or a dict-of-columns (what pandas itself produces).
    Raises a clear error rather than a cryptic pandas traceback if the
    JSON doesn't look like tabular transaction data at all."""
    import json
    try:
        parsed = json.loads(file_bytes.decode("utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError) as e:
        raise ValueError(f"Could not parse this file as JSON: {e}")

    if isinstance(parsed, dict):
        for key in ("transactions", "data", "records", "results"):
            if key in parsed and isinstance(parsed[key], list):
                parsed = parsed[key]
                break

    if isinstance(parsed, list):
        if not parsed:
            raise ValueError("The uploaded JSON file contains an empty list — no transactions to analyze.")
        df = pd.json_normalize(parsed)
    elif isinstance(parsed, dict):
        df = pd.DataFrame(parsed)  # dict-of-columns shape
    else:
        raise ValueError(
            "Uploaded JSON isn't shaped like transaction data — expected a list of "
            "transaction objects (optionally under a 'transactions' or 'data' key)."
        )
    return df


def _read_docx_table(file_bytes):
    """Reads transaction rows out of a Word document. Statements people
    actually paste into Word almost always keep the data as a table, so we
    pull every table on the document, pick the one that looks the most
    like a transaction table (most columns, since a narrow 2-column table
    is more likely a letterhead/address block), and use its first row as
    the header. Falls back to a clear error if the document has no table
    at all — a free-text paragraph-per-transaction layout isn't something
    we can reliably parse."""
    try:
        import docx
    except ImportError:
        raise ValueError(
            "Reading .docx files requires the 'python-docx' package — "
            "install it with: pip install python-docx"
        )
    document = docx.Document(io.BytesIO(file_bytes))
    if not document.tables:
        raise ValueError(
            "Couldn't find a table in this Word document. Transaction uploads need "
            "the data in a table (Insert > Table in Word) with a header row — plain "
            "paragraph text can't be reliably parsed."
        )
    best_table = max(document.tables, key=lambda t: len(t.columns))
    rows = [[cell.text.strip() for cell in row.cells] for row in best_table.rows]
    header, data_rows = rows[0], rows[1:]
    return pd.DataFrame(data_rows, columns=header)


def _validate_and_fill(df, user_id):
    missing = REQUIRED_COLS - set(df.columns)
    if missing:
        raise ValueError(
            f"Uploaded file is missing required column(s): {', '.join(missing)}. "
            f"Columns found in your file: {', '.join(df.columns)}. "
            f"Rename the relevant column(s) to 'amount' and 'merchant' and re-upload."
        )

    if "transaction_id" not in df.columns:
        df["transaction_id"] = [str(uuid.uuid4()) for _ in range(len(df))]
    if "user_id" not in df.columns:
        df["user_id"] = user_id
    if "timestamp" not in df.columns:
        raise ValueError("Uploaded file needs a 'timestamp' or 'date' column.")
    df["timestamp"] = pd.to_datetime(df["timestamp"], errors="coerce")
    if df["timestamp"].isna().all():
        raise ValueError("Could not parse any dates in the timestamp/date column.")

    # Real exports frequently store amount as a currency-formatted string
    # (₹1,450.00, Rs. 820, a lone "+" on a malformed credit row, etc.) —
    # clean it to a plain float before anything downstream touches it.
    df["amount"] = df["amount"].apply(_parse_amount)

    df["category"] = df.get("category", pd.Series(["Other"] * len(df))).fillna("Other")
    df["category"] = fill_categories(df["merchant"], df["category"])
    df["payment_mode"] = df.get("payment_mode", pd.Series(["UPI-Intent"] * len(df))).fillna("UPI-Intent")
    df["recurring_flag"] = 0
    df["subscription_id"] = ""
    df["anomaly_label"] = 0  # unknown for real uploads; ground truth not available

    # Incoming money (a "Credit" in a Debit/Credit 'type' column) isn't
    # spending — counting refunds/transfers received would inflate totals and
    # skew the personal baseline. Same rule the PDF extractor already applies.
    for col in ("type", "txn_type"):
        if col in df.columns:
            is_credit = df[col].astype(str).str.contains("credit", case=False, na=False)
            df = df[~is_credit]
            break

    df = df[df["amount"] > 0].dropna(subset=["timestamp"])
    return df


def _add_date_parts(df):
    df["hour"] = df["timestamp"].dt.hour
    df["day"] = df["timestamp"].dt.day
    df["day_of_week"] = df["timestamp"].dt.dayofweek
    df["month"] = df["timestamp"].dt.month
    df["is_weekend"] = df["day_of_week"].isin([5, 6]).astype(int)
    return df


def _score_anomalies(df, reuse_model=None, reuse_scaler=None):
    """Personal model if enough history, else the trained global model.
    If reuse_model/reuse_scaler are given (the backend's drift-aware flow
    decided NOT to retrain this time), score with those instead of fitting
    a fresh model — this is what makes "skip retrain when nothing drifted"
    actually save compute rather than retraining on every single call."""
    X = df[FEATURE_COLS].fillna(0).values

    if reuse_model is not None and reuse_scaler is not None:
        model, scaler = reuse_model, reuse_scaler
        X_scaled = scaler.transform(X)
        model_used = "personal (reused, no drift since last train)"
    elif len(df) >= MIN_TXNS_FOR_PERSONAL_MODEL:
        scaler = StandardScaler()
        X_scaled = scaler.fit_transform(X)
        model = IsolationForest(n_estimators=100, contamination=0.02, random_state=42, n_jobs=-1)
        model.fit(X_scaled)
        model_used = "personal"
    else:
        model = joblib.load(MODEL_PATH)
        scaler = joblib.load(SCALER_PATH)
        X_scaled = scaler.transform(X)
        model_used = "global (not enough history yet for a personal model)"

    raw = model.decision_function(X_scaled)
    df["anomaly_score"] = 1 - (raw - raw.min()) / (raw.max() - raw.min() + 1e-9)

    # Isolation Forest cannot extrapolate: an amount far BEYOND anything in the
    # training range lands in the same leaf as the largest training amount, so
    # a 20x spike can score as "normal". Add a robust per-user rule on top
    # (median/MAD of log amount — a few spikes can't drag it) so genuinely
    # extreme spending is always caught. Only the HIGH side counts as unusual
    # spending: a tiny purchase is not something to alert on.
    amounts = df["amount"].values.astype(float)
    la = np.log1p(amounts)
    med = np.median(la)
    mad = np.median(np.abs(la - med))
    typical = float(np.expm1(med))
    high_rule = (0.6745 * (la - med) / mad > RULE_MOD_Z) if mad > 0 else np.zeros(len(df), dtype=bool)
    # with very little history the model is shaky, so only call something
    # unusual when it is clearly bigger than the person's normal payment
    min_ratio = 3.0 if len(df) < 40 else 1.0
    forest_flag = (model.predict(X_scaled) == -1) & (amounts > typical * min_ratio)
    # A payment that is ordinary FOR THAT SHOP is not unusual, even if it is
    # large next to everything else (a Rs 2,000 Amazon order when Amazon is
    # normally Rs 700-2,200). The shop's "usual" is its median over payments
    # that are not already extreme for this person, so a spike concentrated
    # in one shop cannot redefine that shop's normal. A shop that charges a
    # steady amount across 3+ different months (rent, fees) is regular too.
    # Needs real history (6+ payments) and never overrides a very extreme amount.
    z = (0.6745 * (la - med) / mad) if mad > 0 else np.zeros(len(df))
    months = df["timestamp"].dt.to_period("M")
    normal_for_shop = np.zeros(len(df), dtype=bool)
    for shop, idx in df.groupby("merchant").indices.items():
        a_shop, z_shop = amounts[idx], z[idx]
        calm = a_shop[z_shop <= RULE_MOD_Z]
        usual = np.median(calm) if len(calm) >= 6 else None
        steady = (len(idx) >= 4 and months.iloc[idx].nunique() >= 3
                  and a_shop.mean() > 0 and a_shop.std() / a_shop.mean() < 0.35)
        if usual is not None:
            normal_for_shop[idx] = a_shop <= 2.5 * usual
        if steady:
            normal_for_shop[idx] = True
    normal_for_shop &= ~(z > 2 * RULE_MOD_Z)
    df["anomaly_prediction"] = ((forest_flag | high_rule) & ~normal_for_shop).astype(int)
    df["anomaly_score"] = np.where(high_rule, np.maximum(df["anomaly_score"], 0.85), df["anomaly_score"])
    df["typical_amount"] = typical
    df["explanation"] = df.apply(explain_row, axis=1)
    df["model_used"] = model_used
    return df


def _spending_breakdown(df):
    by_category = (
        df.groupby("category")["amount"]
        .agg(total="sum", count="count", average="mean")
        .round(2).sort_values("total", ascending=False).reset_index()
        .to_dict(orient="records")
    )
    by_merchant = (
        df.groupby("merchant")["amount"]
        .agg(total="sum", count="count", average="mean")
        .round(2).sort_values("total", ascending=False).head(15).reset_index()
        .to_dict(orient="records")
    )
    by_month = (
        df.groupby(df["timestamp"].dt.to_period("M"))["amount"]
        .sum().round(2)
        .rename_axis("month").reset_index()
    )
    by_month["month"] = by_month["month"].astype(str)
    return {
        "by_category": by_category,
        "by_merchant_top15": by_merchant,
        "by_month": by_month.to_dict(orient="records"),
    }


def _safe_credits(df):
    try:
        from .credit_service import credits_from_frame
        return credits_from_frame(df)
    except Exception:
        return None


def parse_and_validate(file_bytes: bytes, user_id: str, filename: str = "upload.csv"):
    """Split out from analyze_upload so the backend can merge a new upload
    into a user's ACCUMULATED history before re-running analysis on the
    full set (rather than re-analyzing only the latest upload in isolation).
    Deliberately format-agnostic — the person uploading shouldn't have to
    care whether their bank/app gave them a CSV, an Excel sheet, a JSON
    export, a Word document, or a PDF statement; filename decides which
    reader runs, and every path converges on the same tabular frame before
    validation."""
    name = filename.lower()
    if name.endswith(".pdf"):
        from .statement_extractor import extract_transactions_from_pdf
        df, source = extract_transactions_from_pdf(file_bytes)
        try:
            from .credit_service import credits_from_pdf
            credits_df = credits_from_pdf(file_bytes)
        except Exception:      # credits are optional; never break the spending upload
            credits_df = None
        # df already has canonical columns (timestamp, amount, merchant,
        # category, payment_mode) from the extractor — skip _normalize_columns
    elif name.endswith((".docx", ".doc")):
        df = _read_docx_table(file_bytes)
        df = _normalize_columns(df)
        df = _reconstruct_wrapped_rows(df)
        credits_df = _safe_credits(df)
    elif name.endswith(".json"):
        df = _read_json_transactions(file_bytes)
        df = _normalize_columns(df)
        credits_df = _safe_credits(df)
    else:
        if name.endswith((".xlsx", ".xls")):
            df = pd.read_excel(io.BytesIO(file_bytes))
        else:
            df = pd.read_csv(io.BytesIO(file_bytes))
        # Embedded newlines inside a wrapped cell (seen in real .xlsx
        # exports, e.g. "2026,\n08:10 PM") would otherwise break date
        # parsing and text reassembly below.
        _clean_cell = lambda v: v.replace("\n", " ").strip() if isinstance(v, str) else v  # noqa: E731
        if hasattr(df, "map"):
            df = df.map(_clean_cell)
        else:  # pandas < 2.1 fallback
            df = df.applymap(_clean_cell)
        df = _normalize_columns(df)
        df = _reconstruct_wrapped_rows(df)
        credits_df = _safe_credits(df)

    df = _validate_and_fill(df, user_id)
    df.attrs["credits"] = credits_df      # money received; handled separately from spending
    return df


def build_features(df):
    """Clean + engineer features (Phases 2-3), split out so the backend can
    build features BEFORE deciding whether to retrain (Phase 12 needs the
    engineered feature matrix, not the raw transaction rows)."""
    if len(df) < 5:
        raise ValueError("Need at least 5 transactions to run analysis.")
    df = df.sort_values("timestamp").reset_index(drop=True)
    df = _add_date_parts(df)
    df = build_user_spending_features(df)
    df = build_recurring_features(df)
    return df


def analyze_df(df, user_id: str = "uploaded_user", reuse_model=None, reuse_scaler=None):
    """Runs anomaly scoring + subscription detection + breakdown on an
    ALREADY feature-engineered DataFrame (see build_features)."""
    df = _score_anomalies(df, reuse_model=reuse_model, reuse_scaler=reuse_scaler)

    subs = detect_subscriptions(df)
    detected_subs = subs[subs["is_subscription"] == 1].copy()
    detected_subs["next_expected_payment"] = detected_subs["next_expected_payment"].astype(str)
    detected_subs["last_txn"] = detected_subs["last_txn"].astype(str)

    anomalies = (
        df[df["anomaly_prediction"] == 1]
        .sort_values("anomaly_score", ascending=False)
        [["transaction_id", "timestamp", "amount", "merchant", "category",
          "anomaly_score", "explanation", "model_used"]]
        .copy()
    )
    anomalies["timestamp"] = anomalies["timestamp"].astype(str)

    breakdown = _spending_breakdown(df)

    period = detect_period(df)
    insights = generate_insights(df, period, detected_subs)

    summary = {
        "total_transactions": int(len(df)),
        "total_spent": round(float(df["amount"].sum()), 2),
        "date_range": {
            "start": str(df["timestamp"].min().date()),
            "end": str(df["timestamp"].max().date()),
        },
        "anomalies_detected": int(df["anomaly_prediction"].sum()),
        "subscriptions_detected": int(len(detected_subs)),
        "subscriptions_monthly_total": round(float(detected_subs["monthly_cost"].sum()), 2) if len(detected_subs) else 0,
        "model_used": df["model_used"].iloc[0],
        "average_spending": round(float(df["average_spending"].iloc[0]), 2),
        "median_spending": round(float(df["median_spending"].iloc[0]), 2),
        "spending_std": round(float(df["spending_std"].iloc[0]), 2),
    }

    return {
        "summary": summary,
        "anomalies": anomalies.to_dict(orient="records"),
        "subscriptions": detected_subs.to_dict(orient="records"),
        "spending_breakdown": breakdown,
        "period_analysis": period,
        "insights": insights,
    }


def analyze_upload(file_bytes: bytes, user_id: str = "uploaded_user", filename: str = "upload.csv"):
    """Stand-alone, single-upload analysis (no DB/accumulation) — used by
    the anonymous /upload/history endpoint for one-off, no-login use."""
    df = parse_and_validate(file_bytes, user_id, filename)
    df = build_features(df)
    return analyze_df(df, user_id)
